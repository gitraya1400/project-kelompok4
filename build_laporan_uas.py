# -*- coding: utf-8 -*-
"""
Generator Laporan UAS Keamanan Sistem Informasi — Kelompok 4
Topik 4: Keamanan Basis Data

Mengikuti Tabel 2 Soal UAS (struktur BAB I-VII + Daftar Pustaka + Lampiran).
Format: A4, Times New Roman 12pt, spasi 1,5.

Jalankan:  python build_laporan_uas.py
Keluaran:  UAS_KSI_Kelompok4.docx
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(BASE, "user_screenshots")
OUT = os.path.join(BASE, "UAS_KSI_Kelompok4.docx")

TODO = "[PERLU DIISI]"


# --------------------------------------------------------------------------
# Helper dasar
# --------------------------------------------------------------------------
def setup(doc):
    """A4, margin 3-3-3-3 cm, Times New Roman 12pt, spasi 1,5."""
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        s.top_margin = s.bottom_margin = Inches(1.18)
        s.left_margin = Inches(1.18)
        s.right_margin = Inches(1.18)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)


def judul_bab(doc, teks):
    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run(teks)
    r.bold = True
    r.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(12)


def sub(doc, teks, level=2):
    p = doc.add_paragraph()
    r = p.add_run(teks)
    r.bold = True
    r.font.size = Pt(12 if level == 2 else 12)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def par(doc, teks, just=True):
    p = doc.add_paragraph(teks)
    if just:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def poin(doc, teks, lvl=0):
    p = doc.add_paragraph(teks, style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.3 + 0.3 * lvl)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def nomor(doc, teks):
    p = doc.add_paragraph(teks, style="List Number")
    p.paragraph_format.left_indent = Inches(0.3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def kode(doc, teks):
    """Blok kode / cuplikan terminal."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(teks)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    return p


def tabel(doc, header, baris, caption=None, lebar=None):
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = c.add_run(caption)
        rc.bold = True
        rc.font.size = Pt(11)
        c.paragraph_format.space_after = Pt(4)
        c.paragraph_format.line_spacing = 1.0

    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = True
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ""
        pp = hdr[i].paragraphs[0]
        rr = pp.add_run(str(h))
        rr.bold = True
        rr.font.size = Pt(10)
        rr.font.name = "Times New Roman"
        pp.paragraph_format.line_spacing = 1.0
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # shading header
        sh = OxmlElement("w:shd")
        sh.set(qn("w:fill"), "D9E2F3")
        hdr[i]._tc.get_or_add_tcPr().append(sh)

    for row in baris:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pp = cells[i].paragraphs[0]
            rr = pp.add_run(str(v))
            rr.font.size = Pt(10)
            rr.font.name = "Times New Roman"
            pp.paragraph_format.line_spacing = 1.0
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def gambar(doc, nama, caption):
    """Sisipkan screenshot asli dari user_screenshots/ bila ada."""
    path = os.path.join(SHOTS, nama)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(5.9))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[{TODO}: sisipkan {nama}]")
        r.italic = True
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = c.add_run(caption)
    rc.font.size = Pt(10)
    rc.italic = True
    c.paragraph_format.line_spacing = 1.0
    c.paragraph_format.space_after = Pt(10)


def catatan_isi(doc, teks):
    """Penanda bagian yang harus diisi tim — merah & miring."""
    p = doc.add_paragraph()
    r = p.add_run(f"{TODO} — {teks}")
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r.font.size = Pt(11)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


# --------------------------------------------------------------------------
# MAIN
# --------------------------------------------------------------------------
class H:
    """Namespace helper agar modul bab dapat memanggil fungsi pembantu."""
    judul_bab = staticmethod(judul_bab)
    sub = staticmethod(sub)
    par = staticmethod(par)
    poin = staticmethod(poin)
    nomor = staticmethod(nomor)
    kode = staticmethod(kode)
    tabel = staticmethod(tabel)
    gambar = staticmethod(gambar)
    catatan_isi = staticmethod(catatan_isi)


def main():
    import _bab_1_3, _bab_4, _bab_5_7, _lampiran

    doc = Document()
    setup(doc)

    _bab_1_3.sampul(doc, H)
    _bab_1_3.bab1(doc, H)
    _bab_1_3.bab2(doc, H)
    _bab_1_3.bab3(doc, H)
    _bab_4.bab4(doc, H)
    _bab_5_7.bab5(doc, H)
    _bab_5_7.bab6(doc, H)
    _bab_5_7.bab7(doc, H)
    _bab_5_7.pustaka(doc, H)
    _lampiran.lampiran(doc, H)

    doc.save(OUT)
    print("Laporan tersimpan: %s" % OUT)
    print("Jumlah paragraf   : %d" % len(doc.paragraphs))
    print("Jumlah tabel      : %d" % len(doc.tables))


if __name__ == "__main__":
    main()
