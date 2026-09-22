import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import datetime

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right.
    value: dict(sz=12, val='single', color='FF0000')
    """
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for border_name, border_props in kwargs.items():
        b_elm = parse_xml(
            f'<w:{border_name} {nsdecls("w")} '
            f'w:val="{border_props.get("val", "single")}" '
            f'w:sz="{border_props.get("sz", "4")}" '
            f'w:space="0" '
            f'w:color="{border_props.get("color", "auto")}"/>'
        )
        tcBorders.append(b_elm)
    tcPr.append(tcBorders)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_border(cell, 
                    left={'val': 'single', 'sz': '18', 'color': '2563EB'},
                    top={'val': 'single', 'sz': '4', 'color': 'E5E7EB'},
                    bottom={'val': 'single', 'sz': '4', 'color': 'E5E7EB'},
                    right={'val': 'single', 'sz': '4', 'color': 'E5E7EB'})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_screenshot_placeholder(doc, fig_number, title, guidance_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=180, bottom=180, left=200, right=200)
    set_cell_border(cell, 
                    left={'val': 'dashed', 'sz': '8', 'color': '94A3B8'},
                    right={'val': 'dashed', 'sz': '8', 'color': '94A3B8'},
                    top={'val': 'dashed', 'sz': '8', 'color': '94A3B8'},
                    bottom={'val': 'dashed', 'sz': '8', 'color': '94A3B8'})
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run("📸 [ TEMPAT MENEMPELKAN SCREENSHOT ]\n")
    r1.font.name = 'Arial'
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(37, 99, 235)
    
    r2 = p.add_run(guidance_text)
    r2.font.name = 'Arial'
    r2.font.italic = True
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(100, 116, 139)
    
    # Empty space for image insertion
    p_space = cell.add_paragraph()
    p_space.paragraph_format.space_before = Pt(24)
    p_space.paragraph_format.space_after = Pt(24)
    p_space.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_space = p_space.add_run("(Tempelkan gambar hasil tangkapan layar di sini)")
    r_space.font.name = 'Arial'
    r_space.font.size = Pt(8.5)
    r_space.font.color.rgb = RGBColor(148, 163, 184)
    
    # Caption below table
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(f"Gambar {fig_number}: {title}")
    r_cap.font.name = 'Arial'
    r_cap.font.bold = True
    r_cap.font.size = Pt(9.5)
    r_cap.font.color.rgb = RGBColor(51, 65, 85)

def add_callout(doc, title, text, box_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    if box_type == "warning":
        bg_color = "FEF2F2"
        border_color = "DC2626"
        title_color = RGBColor(185, 28, 28)
    elif box_type == "success":
        bg_color = "F0FDF4"
        border_color = "16A34A"
        title_color = RGBColor(21, 128, 61)
    else: # info
        bg_color = "EFF6FF"
        border_color = "2563EB"
        title_color = RGBColor(29, 78, 216)
        
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_border(cell, 
                    left={'val': 'single', 'sz': '24', 'color': border_color},
                    top={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                    bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                    right={'val': 'single', 'sz': '4', 'color': 'E2E8F0'})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.font.name = 'Arial'
    r_title.font.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = title_color
    
    r_body = p.add_run(text)
    r_body.font.name = 'Arial'
    r_body.font.size = Pt(9)
    r_body.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Helper functions defined successfully.")
