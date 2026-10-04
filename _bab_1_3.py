# -*- coding: utf-8 -*-
"""Bagian Sampul, BAB I, II, III — diimpor oleh build_laporan_uas.py"""
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def sampul(doc, H):
    for _ in range(3):
        doc.add_paragraph()
    for teks, sz in [("LAPORAN PROYEK AKHIR", 16),
                     ("MATA KULIAH KEAMANAN SISTEM INFORMASI", 14)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(teks); r.bold = True; r.font.size = Pt(sz)

    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IMPLEMENTASI HARDENING, KONTROL AKSES BERBASIS PERAN, "
                  "ENKRIPSI, AUDIT LOG, DAN HIGH AVAILABILITY PADA BASIS DATA "
                  "SISTEM INFORMASI KLINIK")
    r.bold = True; r.font.size = Pt(14)

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Disusun oleh:").font.size = Pt(12)
    doc.add_paragraph()

    for i in range(1, 6):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.add_run("[Nama Lengkap Anggota %d]" % i).font.size = Pt(12)
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(8)
        p2.add_run("[NIM Anggota %d]" % i).font.size = Pt(12)

    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Dosen Pengampu:").font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Farid Ridho"); r.bold = True; r.font.size = Pt(12)

    for _ in range(3):
        doc.add_paragraph()
    for teks in ["PROGRAM STUDI KOMPUTASI STATISTIK",
                 "POLITEKNIK STATISTIKA STIS", "2026"]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(teks); r.bold = True; r.font.size = Pt(12)
        p.paragraph_format.space_after = Pt(2)


def bab1(doc, H):
    H.judul_bab(doc, "BAB I  PENDAHULUAN")

    H.sub(doc, "1.1  Latar Belakang")
    H.par(doc, "Badan Pusat Statistik dan institusi penyelenggara statistik lainnya mengelola "
          "data individu dalam jumlah besar, mulai dari data kependudukan, ketenagakerjaan, "
          "hingga data kesehatan. Data semacam ini tergolong data pribadi spesifik menurut "
          "Undang-Undang Nomor 27 Tahun 2022 tentang Pelindungan Data Pribadi, dan sekaligus "
          "dilindungi prinsip kerahasiaan sebagaimana diatur Undang-Undang Nomor 16 Tahun 1997 "
          "tentang Statistik. Kebocoran satu baris data responden tidak hanya merugikan individu, "
          "tetapi juga meruntuhkan kepercayaan publik terhadap lembaga statistik, yaitu modal "
          "utama yang menentukan tingkat respons survei pada periode berikutnya.")
    H.par(doc, "Basis data merupakan titik paling kritis dalam rantai pengolahan data karena "
          "menjadi muara seluruh informasi. Namun konfigurasi bawaan sistem manajemen basis data "
          "umumnya mengutamakan kemudahan pemasangan, bukan keamanan. MySQL pada konfigurasi "
          "standar tidak mewajibkan enkripsi lalu lintas jaringan, tidak mengaktifkan pencatatan "
          "audit ke tabel, dan membiarkan administrator memberikan hak akses melebihi kebutuhan. "
          "OWASP secara konsisten menempatkan Injection dan Security Misconfiguration dalam "
          "sepuluh besar risiko aplikasi web, dan keduanya bermuara pada lapisan basis data [1].")
    H.par(doc, "Selain aspek kerahasiaan, aspek ketersediaan juga kerap terabaikan. Arsitektur "
          "basis data tunggal menjadikan server sebagai single point of failure: ketika server "
          "berhenti, seluruh layanan ikut berhenti. Pada konteks sistem informasi statistik, "
          "terhentinya layanan saat periode pengumpulan data berpotensi menyebabkan kehilangan "
          "data lapangan yang tidak dapat dipulihkan.")
    H.par(doc, "Proyek ini membangun lingkungan laboratorium terisolasi berisi klaster basis data "
          "klinik bernama klinik_db yang mengintegrasikan pengamanan menyeluruh, meliputi enkripsi "
          "transport dan penyimpanan, kontrol akses berbasis peran, pencatatan audit, serta "
          "replikasi dengan failover otomatis. Efektivitas setiap kontrol kemudian dibuktikan "
          "melalui lima skenario pengujian yang terukur dan dapat direproduksi.")

    H.sub(doc, "1.2  Rumusan Masalah")
    for t in [
        "Bagaimana menerapkan hardening pada sistem manajemen basis data MySQL sehingga memenuhi "
        "aspek kerahasiaan, integritas, dan ketersediaan sesuai kerangka CIA Triad?",
        "Sejauh mana kontrol keamanan yang diterapkan, yaitu prepared statement, least privilege, "
        "TLS, enkripsi kolom, dan audit log, efektif menahan serangan SQL injection, eskalasi hak "
        "akses, serta intersepsi koneksi?",
        "Bagaimana arsitektur replikasi master-slave dengan load balancer menjaga ketersediaan "
        "layanan ketika server utama mengalami kegagalan?",
        "Bagaimana rancangan pengelolaan keamanan data dan layanan teknologi informasi yang sesuai "
        "UU PDP, UU Statistik, NIST SP 800-61, dan ISO/IEC 27001:2022 untuk konteks sistem "
        "informasi statistik?",
    ]:
        H.nomor(doc, t)

    H.sub(doc, "1.3  Tujuan Proyek")
    for t in [
        "Membangun klaster basis data terhardening pada lingkungan laboratorium terisolasi dengan "
        "lima kontrol keamanan aktif dan terverifikasi.",
        "Membuktikan efektivitas setiap kontrol melalui lima skenario pengujian berpola "
        "Before-Attack, During-Attack, dan After-Mitigation dengan parameter keberhasilan terukur.",
        "Mengidentifikasi minimal tiga temuan kerentanan beserta tingkat risiko CVSS v3.1, akar "
        "penyebab, tindakan perbaikan, dan hasil uji ulang.",
        "Menyusun rancangan pengelolaan keamanan data dan layanan teknologi informasi yang mencakup "
        "pelindungan data pribadi, penanganan insiden, audit keamanan, dan peta jalan implementasi.",
    ]:
        H.nomor(doc, t)

    H.sub(doc, "1.4  Manfaat Proyek")
    H.par(doc, "Secara akademis, proyek ini menyediakan studi kasus terdokumentasi mengenai "
          "penerapan CIA Triad pada lapisan basis data, lengkap dengan bukti eksekusi yang dapat "
          "direproduksi pembaca lain melalui berkas konfigurasi yang disertakan pada lampiran.")
    H.par(doc, "Secara praktis, seluruh konfigurasi, skrip pengujian, dan prosedur operasional yang "
          "dihasilkan dapat diadaptasi langsung oleh pengelola sistem informasi statistik yang "
          "menggunakan MySQL, khususnya pada aspek penegakan TLS, pemberian hak akses minimum, "
          "dan pembangunan jejak audit untuk keperluan forensik.")

    H.sub(doc, "1.5  Ruang Lingkup dan Batasan")
    H.par(doc, "Ruang lingkup proyek meliputi:")
    for t in [
        "Basis data klinik_db pada MySQL 8.0.44 dengan empat tabel, yaitu users, pasien, "
        "rekam_medis, dan audit_log.",
        "Empat vektor ancaman yang diuji: SQL injection, eskalasi hak akses, intersepsi koneksi "
        "tanpa enkripsi, dan kegagalan server tunggal.",
        "Arsitektur dua node MySQL berpola master-slave dengan HAProxy sebagai load balancer.",
        "Lingkungan laboratorium terisolasi berbasis Docker Compose pada jaringan bridge privat "
        "tanpa akses dari jaringan luar.",
    ]:
        H.poin(doc, t)

    H.par(doc, "Batasan proyek meliputi:")
    for t in [
        "Implementasi difokuskan pada satu jalur, yaitu Docker. Jalur alternatif berbasis XAMPP "
        "dan Laragon dikeluarkan dari ruang lingkup karena versi MySQL serta lokasi berkas "
        "konfigurasinya bergantung pada instalasi tiap mesin, sehingga kesetaraan konfigurasi "
        "antarjalur tidak dapat dijamin maupun diverifikasi.",
        "Sertifikat TLS yang digunakan bersifat self-signed dari Certificate Authority lokal, "
        "bukan dari CA publik tepercaya.",
        "Kunci enkripsi AES masih ditanamkan dalam skrip untuk keperluan demonstrasi laboratorium. "
        "Implementasi produksi memerlukan Key Management System terpisah.",
        "HAProxy dijalankan sebagai instans tunggal sehingga masih berperan sebagai single point "
        "of failure pada lapisan load balancer.",
        "Data yang digunakan merupakan data sintetis, bukan data pasien sesungguhnya.",
    ]:
        H.poin(doc, t)


def bab2(doc, H):
    H.judul_bab(doc, "BAB II  TINJAUAN PUSTAKA")

    H.sub(doc, "2.1  Landasan Teori")
    H.sub(doc, "2.1.1  CIA Triad", 3)
    H.par(doc, "CIA Triad merupakan kerangka dasar keamanan informasi yang terdiri atas tiga pilar. "
          "Confidentiality menjamin data hanya dapat diakses pihak berwenang. Integrity menjamin "
          "data tidak diubah tanpa otorisasi. Availability menjamin data dan layanan tersedia saat "
          "dibutuhkan [2]. Pada praktik pengamanan basis data, ketiga pilar ini sering dilengkapi "
          "pilar keempat, yaitu Accountability, yang menjamin setiap tindakan dapat ditelusuri "
          "kepada pelakunya sehingga mendukung non-repudiation.")
    H.par(doc, "Proyek ini memetakan kelima skenario pengujian ke pilar-pilar tersebut: SQL "
          "injection dan enkripsi menguji Confidentiality dan Integrity, least privilege menguji "
          "Confidentiality dan Integrity, failover menguji Availability, sedangkan audit logging "
          "menguji Accountability.")

    H.sub(doc, "2.1.2  Standar dan Kerangka Kerja Keamanan", 3)
    H.par(doc, "ISO/IEC 27001:2022 menetapkan persyaratan sistem manajemen keamanan informasi "
          "beserta 93 kontrol pada Annex A yang terbagi menjadi empat tema: organisasional, orang, "
          "fisik, dan teknologi [3]. Kontrol yang paling relevan dengan proyek ini mencakup A.8.3 "
          "Information Access Restriction, A.8.24 Use of Cryptography, dan A.8.15 Logging.")
    H.par(doc, "NIST Special Publication 800-61 Revision 2 menyediakan panduan penanganan insiden "
          "keamanan komputer dengan empat fase siklus: Preparation; Detection and Analysis; "
          "Containment, Eradication, and Recovery; serta Post-Incident Activity [4]. Panduan ini "
          "menjadi dasar penyusunan rencana tanggap insiden pada Bab V.")
    H.par(doc, "OWASP Top 10 menempatkan Injection pada peringkat A03:2021 dan Security "
          "Misconfiguration pada A05:2021 [1]. CIS Benchmark for MySQL memberikan rekomendasi "
          "konfigurasi teknis terperinci, antara lain penegakan koneksi terenkripsi dan pembatasan "
          "hak akses pengguna aplikasi [5]. NIST SP 800-111 memberikan panduan perlindungan data "
          "tersimpan melalui enkripsi [6].")

    H.sub(doc, "2.1.3  Teknik Keamanan Basis Data", 3)
    H.par(doc, "Prepared statement memisahkan struktur kueri dari data masukan. Struktur kueri "
          "dikirim terlebih dahulu ke server untuk dikompilasi, kemudian nilai parameter dikirim "
          "terpisah sehingga tidak pernah dievaluasi sebagai perintah SQL. Mekanisme inilah yang "
          "menjadikan prepared statement sebagai mitigasi paling efektif terhadap SQL injection.")
    H.par(doc, "Prinsip least privilege mensyaratkan setiap akun hanya memperoleh hak akses minimum "
          "yang diperlukan untuk menjalankan fungsinya. Penerapan prinsip ini membatasi dampak "
          "ketika kredensial satu akun bocor, karena kerusakan terbatas pada cakupan hak akun "
          "tersebut.")
    H.par(doc, "Transport Layer Security melindungi data in transit melalui enkripsi kanal "
          "komunikasi, sedangkan enkripsi pada level kolom melindungi data at rest sehingga data "
          "tetap tidak terbaca apabila media penyimpanan fisik berpindah tangan. Replikasi "
          "master-slave yang dipadukan load balancer menjaga ketersediaan layanan dengan "
          "mengalihkan lalu lintas secara otomatis ketika node utama tidak merespons.")

    H.sub(doc, "2.2  Penelitian dan Proyek Terkait")
    H.tabel(doc,
        ["No", "Referensi", "Fokus", "Perbandingan dengan Proyek Ini"],
        [
            ["1", "OWASP Top 10:2021 [1]",
             "Klasifikasi sepuluh risiko teratas aplikasi web, termasuk A03 Injection.",
             "Menjadi acuan pemilihan vektor uji. Proyek ini memperluasnya ke lapisan basis data "
             "dengan bukti eksekusi before-after, bukan sekadar klasifikasi risiko."],
            ["2", "CIS MySQL Benchmark [5]",
             "Rekomendasi konfigurasi aman MySQL secara terperinci.",
             "Dipakai sebagai dasar hardening. Proyek ini menambahkan pengujian empiris atas "
             "efektivitas tiap rekomendasi, bukan hanya penerapan daftar periksa."],
            ["3", "NIST SP 800-111 [6]",
             "Panduan enkripsi data tersimpan pada perangkat pengguna akhir.",
             "Prinsipnya diterapkan pada level kolom basis data melalui AES_ENCRYPT dengan "
             "pengujian dekripsi menggunakan kunci benar dan kunci salah."],
            ["4", "NIST SP 800-61 Rev.2 [4]",
             "Kerangka penanganan insiden keamanan komputer.",
             "Diadopsi pada Bab V sebagai dasar playbook dan simulasi tabletop yang disesuaikan "
             "dengan konteks kebocoran data pasien."],
            ["5", "ISO/IEC 27001:2022 [3]",
             "Persyaratan SMKI dan 93 kontrol Annex A.",
             "Dipakai sebagai kriteria audit pada Bab V.3 dan kerangka pemetaan kontrol pada "
             "Bab V.4."],
        ],
        caption="Tabel 2.1 Perbandingan Penelitian dan Proyek Terkait")

    H.par(doc, "Perbedaan utama proyek ini dengan kelima referensi tersebut terletak pada "
          "integrasinya. Referensi yang ada umumnya membahas satu aspek secara mendalam, sedangkan "
          "proyek ini menggabungkan hardening, kontrol akses, kriptografi, audit, dan high "
          "availability dalam satu arsitektur, kemudian membuktikan seluruhnya melalui pengujian "
          "yang dapat direproduksi dari satu berkas konfigurasi.")

    H.sub(doc, "2.3  Teknologi dan Tools yang Digunakan")
    H.tabel(doc,
        ["Teknologi", "Versi", "Fungsi", "Alasan Pemilihan"],
        [
            ["MySQL", "8.0.44", "Sistem manajemen basis data target",
             "Mendukung TLS bawaan, GTID replication, general log ke tabel, dan AES_ENCRYPT "
             "tanpa komponen tambahan."],
            ["HAProxy", "2.8", "Load balancer dan health check",
             "Ringan, mendukung mode TCP untuk protokol MySQL, dan menyediakan dashboard "
             "pemantauan bawaan."],
            ["Docker Compose", "v2", "Orkestrasi lingkungan laboratorium",
             "Menjamin reproducibility: seluruh lingkungan terdefinisi dalam satu berkas dan "
             "terisolasi pada jaringan bridge privat."],
            ["OpenSSL", "3.x", "Pembuatan CA, sertifikat server dan klien",
             "Standar de facto pembuatan sertifikat X.509 dan tersedia lintas platform."],
            ["Python", "3.13", "Backend aplikasi peraga dan generator laporan",
             "Pustaka standar memadai untuk HTTP server sederhana tanpa dependensi berat."],
        ],
        caption="Tabel 2.2 Teknologi dan Tools yang Digunakan")


def bab3(doc, H):
    H.judul_bab(doc, "BAB III  METODOLOGI")

    H.sub(doc, "3.1  Tahapan Pelaksanaan")
    H.par(doc, "Pelaksanaan proyek mengikuti lima tahap berurutan. Setiap tahap menghasilkan "
          "keluaran yang menjadi masukan tahap berikutnya.")
    H.tabel(doc,
        ["Tahap", "Kegiatan", "Keluaran"],
        [
            ["1. Studi literatur dan analisis kebutuhan",
             "Kajian CIA Triad, OWASP Top 10, CIS MySQL Benchmark, NIST SP 800-111; identifikasi "
             "empat vektor ancaman; penetapan metrik keberhasilan.",
             "Dokumen analisis kebutuhan dan daftar referensi."],
            ["2. Perancangan arsitektur",
             "Perancangan topologi master-slave dengan HAProxy; desain skema empat tabel; "
             "perancangan RBAC tiga peran; penentuan skema enkripsi.",
             "Diagram arsitektur, skema basis data, spesifikasi skenario uji."],
            ["3. Implementasi dan konfigurasi",
             "Penyusunan docker-compose.yml, pembuatan sertifikat OpenSSL, penerapan GRANT, "
             "konfigurasi replikasi GTID, pengaktifan audit log.",
             "Lingkungan laboratorium berfungsi penuh."],
            ["4. Pengujian dan evaluasi",
             "Eksekusi lima skenario berpola before-during-after; pencatatan bukti; analisis "
             "temuan dan tingkat risiko; uji ulang setelah perbaikan.",
             "Laporan hasil pengujian beserta bukti eksekusi."],
            ["5. Dokumentasi dan penyusunan laporan",
             "Penyusunan laporan, rancangan pengelolaan keamanan, dan materi presentasi.",
             "Laporan akhir dan bahan presentasi."],
        ],
        caption="Tabel 3.1 Tahapan Pelaksanaan Proyek")

    H.sub(doc, "3.2  Arsitektur dan Desain Sistem")
    H.par(doc, "Lingkungan pengujian dibangun pada jaringan bridge privat Docker bernama dbnet "
          "yang terisolasi dari jaringan luar. Seluruh akses klien diarahkan melalui HAProxy pada "
          "porta 3300, sedangkan akses langsung ke node basis data hanya tersedia untuk keperluan "
          "verifikasi. Gambar 3.1 menyajikan topologi lengkap.")
    H.kode(doc,
        "                 Klien / Skrip Demo\n"
        "                          |\n"
        "                          v\n"
        "          +-------------------------------+\n"
        "          |   HAProxy  (porta 3300)       |\n"
        "          |   dashboard  : porta 8900     |\n"
        "          |   health check tiap 2 detik   |\n"
        "          +---------------+---------------+\n"
        "                          |\n"
        "            +-------------+-------------+\n"
        "            v                           v\n"
        "      db-master                     db-slave\n"
        "      server-id = 1                 server-id = 2\n"
        "      host porta 3305               host porta 3307\n"
        "      (menerima tulis)              (read_only + super_read_only)\n"
        "            |                           ^\n"
        "            +--- replikasi GTID over SSL +\n"
        "                 (binlog_format = ROW)")
    H.par(doc, "Gambar 3.1 Topologi jaringan laboratorium terisolasi", just=False)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    H.par(doc, "Node slave ditandai sebagai backup pada konfigurasi HAProxy. Dengan demikian "
          "seluruh lalu lintas diarahkan ke master selama master sehat, dan baru dialihkan ke "
          "slave ketika health check master gagal. Health check menggunakan mode tcp-check yang "
          "hanya membuka soket TCP tanpa mencapai lapisan autentikasi, sehingga tetap berfungsi "
          "meskipun kebijakan TLS diwajibkan.")

    H.par(doc, "Skema basis data klinik_db terdiri atas empat tabel sebagaimana disajikan "
          "Tabel 3.2.")
    H.tabel(doc,
        ["Tabel", "Isi", "Kontrol Keamanan yang Melekat"],
        [
            ["users", "Akun internal: admin, dokter, resepsionis",
             "Menyimpan hash kata sandi; hanya dapat diakses root"],
            ["pasien", "Identitas pasien termasuk NIK",
             "Kolom nik_encrypted bertipe VARBINARY, dienkripsi AES_ENCRYPT"],
            ["rekam_medis", "Diagnosa dan resep",
             "Foreign key ke pasien dan users; akses dibatasi app_user"],
            ["audit_log", "Jejak percobaan mencurigakan",
             "Diisi otomatis oleh prosedur sp_harvest_audit dari mysql.general_log"],
        ],
        caption="Tabel 3.2 Skema Basis Data klinik_db")

    H.par(doc, "Kontrol akses dirancang mengikuti prinsip least privilege dengan tiga peran "
          "sebagaimana disajikan Tabel 3.3. Seluruh akun diwajibkan menggunakan koneksi "
          "terenkripsi melalui klausa REQUIRE SSL, dan dilengkapi pembatasan laju kueri untuk "
          "mengurangi risiko penyalahgunaan.")
    H.tabel(doc,
        ["Peran", "Hak Akses", "Pembatasan Laju"],
        [
            ["app_user", "SELECT, INSERT, UPDATE pada pasien dan rekam_medis; INSERT pada audit_log",
             "1000 kueri/jam, 100 koneksi/jam, 500 pembaruan/jam"],
            ["read_only", "SELECT pada pasien saja",
             "500 kueri/jam, 50 koneksi/jam, 250 pembaruan/jam"],
            ["replicator", "REPLICATION SLAVE pada seluruh basis data; tanpa akses baca data",
             "1000 kueri/jam, 100 koneksi/jam, 500 pembaruan/jam"],
        ],
        caption="Tabel 3.3 Rancangan Kontrol Akses Berbasis Peran")

    H.sub(doc, "3.3  Skenario Pengujian")
    H.par(doc, "Pengujian dilaksanakan melalui lima skenario yang masing-masing mengikuti pola "
          "Before-Attack, During-Attack, dan After-Mitigation. Tabel 3.4 menyajikan skenario "
          "beserta parameter keberhasilan dan metrik evaluasinya.")
    H.tabel(doc,
        ["No", "Skenario", "Aspek", "Parameter Keberhasilan", "Metrik"],
        [
            ["1", "Failover dan High Availability", "Availability",
             "Slave tetap melayani permintaan baca saat master dimatikan; replikasi tersambung "
             "kembali tanpa intervensi manual",
             "Nilai @@server_id yang melayani; status Replica_IO_Running dan Replica_SQL_Running; "
             "Seconds_Behind_Source"],
            ["2", "SQL Injection", "Confidentiality, Integrity",
             "Kueri rentan membocorkan seluruh baris; kueri dengan prepared statement "
             "mengembalikan nol baris pada payload identik",
             "Jumlah baris terekspos pada kondisi rentan dan termitigasi"],
            ["3", "SSL/TLS dan Enkripsi Data", "Confidentiality",
             "Koneksi tanpa enkripsi ditolak; koneksi bersertifikat berhasil; dekripsi dengan "
             "kunci salah menghasilkan NULL",
             "Kode galat penolakan; nama cipher aktif; hasil AES_DECRYPT"],
            ["4", "Least Privilege", "Confidentiality, Integrity",
             "Setiap operasi di luar kewenangan ditolak sistem",
             "Kode galat penolakan pada tiap kombinasi peran dan operasi"],
            ["5", "Audit Logging", "Accountability",
             "Seluruh percobaan dari skenario lain tercatat lengkap dengan waktu, jenis aksi, "
             "tabel sasaran, dan alamat IP",
             "Jumlah dan jenis entri pada tabel audit_log"],
        ],
        caption="Tabel 3.4 Skenario Pengujian, Parameter Keberhasilan, dan Metrik Evaluasi")

    H.sub(doc, "3.4  Etika Pengujian dan Rules of Engagement")
    H.par(doc, "Seluruh pengujian dilaksanakan mengikuti aturan keterlibatan yang disepakati "
          "sebelum eksekusi dimulai. Dokumen lengkap disertakan pada Lampiran A.")
    H.tabel(doc,
        ["Aspek", "Ketentuan"],
        [
            ["Target", "Container db-master, db-slave, dan haproxy pada jaringan bridge privat "
                       "dbnet. Tidak ada sistem pihak ketiga, sistem produksi, maupun layanan "
                       "publik yang menjadi sasaran."],
            ["Isolasi", "Jaringan bridge Docker tanpa port forwarding ke jaringan kampus maupun "
                        "internet. Porta yang dipublikasikan hanya terikat pada localhost mesin "
                        "penguji."],
            ["Data", "Seluruh data merupakan data sintetis yang dibangkitkan tim. Tidak ada data "
                     "pasien, responden, atau data pribadi nyata yang digunakan."],
            ["Metode yang diizinkan", "SQL injection pada basis data milik sendiri, percobaan "
                                      "eskalasi hak akses dengan kredensial yang dibuat tim, "
                                      "penghentian container secara terkendali, dan koneksi tanpa "
                                      "enkripsi untuk menguji penolakan."],
            ["Metode yang dilarang", "Denial of service, eksploitasi terhadap host, pemindaian "
                                     "jaringan di luar dbnet, penggunaan sampel malware biner, "
                                     "serta setiap tindakan yang menyentuh sistem di luar "
                                     "laboratorium."],
            ["Jadwal", "1 sampai dengan 5 Oktober 2026, pada mesin pengembangan anggota tim."],
            ["Penanggung jawab", "[PERLU DIISI: nama ketua kelompok] selaku koordinator pengujian; "
                                 "seluruh anggota menandatangani Pernyataan Etika pada Lampiran B."],
            ["Penanganan temuan", "Temuan dicatat pada laporan ini dan tidak dipublikasikan di "
                                  "luar lingkungan akademik mata kuliah."],
        ],
        caption="Tabel 3.5 Rules of Engagement Pengujian")
